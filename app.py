"""Single-page Streamlit interface for the plate map reformatter."""

from __future__ import annotations

import html
import re
import tempfile
import uuid
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import streamlit as st

from plate_reformatter import (
    complete_source_entry, find_duplicate_lookup_groups, inspect_workbook,
    is_numeric_value, reformat, source_entry_issues,
)


def _excel_color(color):
    """Return an HTML-safe RGB color for common Excel colors."""
    if color is None:
        return None
    if color.type == "rgb" and color.rgb:
        return f"#{color.rgb[-6:]}"
    if color.type == "indexed" and color.indexed is not None:
        try:
            from openpyxl.styles.colors import COLOR_INDEX

            return f"#{COLOR_INDEX[color.indexed][-6:]}"
        except (IndexError, TypeError):
            return None
    return None


def _excel_border(side):
    """Convert a saved Excel border side to a small CSS border declaration."""
    if side is None or side.style is None:
        return None
    css_style = {
        "dashDot": "dashed", "dashDotDot": "dashed", "dashed": "dashed",
        "dotted": "dotted", "double": "double", "hair": "solid",
        "medium": "solid", "mediumDashDot": "dashed",
        "mediumDashDotDot": "dashed", "mediumDashed": "dashed",
        "slantDashDot": "dashed", "thin": "solid", "thick": "solid",
    }.get(side.style, "solid")
    css_width = "2px" if side.style in {"medium", "mediumDashed", "mediumDashDot", "mediumDashDotDot"} else "3px" if side.style == "thick" else "1px"
    return f"{css_width} {css_style} {_excel_color(side.color) or '#000000'}"


def _worksheet_preview_html(worksheet):
    """Render worksheet values and common cell formatting as an Excel-like grid."""
    max_row = max(worksheet.max_row, 1)
    max_column = max(worksheet.max_column, 1)
    column_names = [get_column_letter(column) for column in range(1, max_column + 1)]
    default_gridlines = worksheet.sheet_view.showGridLines is not False
    merged_anchors = {}
    merged_covered = set()
    for merged_range in worksheet.merged_cells.ranges:
        merged_anchors[(merged_range.min_row, merged_range.min_col)] = (
            merged_range.max_row - merged_range.min_row + 1,
            merged_range.max_col - merged_range.min_col + 1,
        )
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for column in range(merged_range.min_col, merged_range.max_col + 1):
                if (row, column) != (merged_range.min_row, merged_range.min_col):
                    merged_covered.add((row, column))

    widths = []
    for column_name in column_names:
        width = worksheet.column_dimensions[column_name].width or 8.43
        widths.append(max(32, min(600, round(width * 7 + 5))))

    parts = [
        '<div class="excel-preview-scroll">',
        '<table class="excel-preview-grid" aria-label="',
        html.escape(f"{worksheet.title} worksheet preview", quote=True),
        '"><colgroup><col style="width:46px">',
    ]
    parts.extend(f'<col style="width:{width}px">' for width in widths)
    parts.append('</colgroup><thead><tr><th class="excel-preview-corner"></th>')
    parts.extend(
        f'<th class="excel-preview-column-heading">{column_name}</th>'
        for column_name in column_names
    )
    parts.append('</tr></thead><tbody>')

    for row_number in range(1, max_row + 1):
        row_height = worksheet.row_dimensions[row_number].height
        row_height_px = max(20, round((row_height or 15) * 4 / 3))
        parts.append(f'<tr style="height:{row_height_px}px"><th class="excel-preview-row-heading">{row_number}</th>')
        for column_number in range(1, max_column + 1):
            if (row_number, column_number) in merged_covered:
                continue
            cell = worksheet.cell(row=row_number, column=column_number)
            styles = ["padding:3px 6px", "min-width:32px", "vertical-align:middle"]
            if default_gridlines:
                styles.append("border:1px solid #d9d9d9")

            fill = cell.fill
            if fill.fill_type == "solid":
                fill_color = _excel_color(fill.fgColor)
                if fill_color:
                    styles.append(f"background-color:{fill_color}")

            font = cell.font
            if font.name:
                family = re.sub(r"[^A-Za-z0-9 ,_-]", "", font.name)
                if family:
                    styles.append(f"font-family:'{family}',Arial,sans-serif")
            if font.sz:
                styles.append(f"font-size:{max(7, min(36, font.sz))}pt")
            if font.bold:
                styles.append("font-weight:bold")
            if font.italic:
                styles.append("font-style:italic")
            decorations = []
            if font.underline:
                decorations.append("underline")
            if font.strike:
                decorations.append("line-through")
            if decorations:
                styles.append(f"text-decoration:{' '.join(decorations)}")
            font_color = _excel_color(font.color)
            if font_color:
                styles.append(f"color:{font_color}")

            alignment = cell.alignment
            if alignment.horizontal in {"left", "center", "right", "justify"}:
                styles.append(f"text-align:{alignment.horizontal}")
            if alignment.vertical in {"top", "center", "bottom"}:
                vertical = "middle" if alignment.vertical == "center" else alignment.vertical
                styles.append(f"vertical-align:{vertical}")
            styles.append("white-space:pre-wrap" if alignment.wrap_text else "white-space:pre")

            border = cell.border
            for side_name, css_name in (
                ("left", "border-left"), ("right", "border-right"),
                ("top", "border-top"), ("bottom", "border-bottom"),
            ):
                declaration = _excel_border(getattr(border, side_name))
                if declaration:
                    styles.append(f"{css_name}:{declaration}")

            row_span, col_span = merged_anchors.get((row_number, column_number), (1, 1))
            span_attributes = ""
            if row_span > 1:
                span_attributes += f' rowspan="{row_span}"'
            if col_span > 1:
                span_attributes += f' colspan="{col_span}"'
            value = "" if cell.value is None else str(cell.value)
            display_value = html.escape(value, quote=False).replace("\n", "<br>") or "&nbsp;"
            parts.append(
                f'<td{span_attributes} style="{";".join(styles)}">{display_value}</td>'
            )
        parts.append('</tr>')
    parts.append('</tbody></table></div>')
    parts.insert(0, """<style>
      .excel-preview-scroll { max-width:100%; max-height:72vh; overflow:auto; border:1px solid #b7b7b7; background:#fff; }
      .excel-preview-grid { border-collapse:collapse; table-layout:fixed; color:#000; background:#fff; font-family:Arial,sans-serif; font-size:11pt; }
      .excel-preview-grid th { position:sticky; z-index:2; padding:3px 5px; border:1px solid #b7b7b7; background:#f3f3f3; color:#333; font:10pt Arial,sans-serif; text-align:center; }
      .excel-preview-grid thead th { top:0; z-index:3; }
      .excel-preview-grid th.excel-preview-row-heading { left:0; z-index:2; min-width:46px; font-weight:normal; }
      .excel-preview-grid th.excel-preview-corner { top:0; left:0; z-index:4; min-width:46px; }
      .excel-preview-grid td { box-sizing:border-box; overflow-wrap:anywhere; }
    </style>""")
    return "".join(parts)


st.set_page_config(page_title="96-Well Plate Map Reformatter", page_icon="🧪")
st.title("96-Well Plate Map Reformatter")
st.write("Upload a plate workbook, enter a plate name, and create a reformatted workbook.")

uploaded_file = st.file_uploader(
    "Source Excel workbook",
    type=["xlsx", "xls"],
    help="Supports modern .xlsx files and legacy .xls files.",
    disabled=st.session_state.get("pending_review") is not None,
)
plate_name = st.text_input(
    "Plate name", placeholder="For example, T01",
    disabled=st.session_state.get("pending_review") is not None,
)

if "pending_review" not in st.session_state:
    st.session_state.pending_review = None
if "completed_result" not in st.session_state:
    st.session_state.completed_result = None

if st.button(
    "Check workbook", type="primary",
    disabled=st.session_state.get("pending_review") is not None,
):
    if uploaded_file is None:
        st.error("Choose an Excel workbook before starting.")
    elif not plate_name.strip():
        st.error("Enter a plate name before starting.")
    else:
        try:
            with st.spinner("Checking workbook for required information…"):
                suffix = Path(uploaded_file.name).suffix.lower()
                with tempfile.TemporaryDirectory(prefix="plate_reformatter_") as work_dir:
                    source_path = Path(work_dir) / f"source{suffix}"
                    source_path.write_bytes(uploaded_file.getvalue())
                    inspection = inspect_workbook(source_path)
            st.session_state.pending_review = {
                "source_bytes": uploaded_file.getvalue(),
                "filename": uploaded_file.name,
                "suffix": suffix,
                "plate_name": plate_name.strip(),
                "missing": inspection["missing"],
                "warnings": inspection["warnings"],
                "study": inspection["study"],
                "review_id": uuid.uuid4().hex,
            }
            st.session_state.completed_result = None
        except Exception as exc:
            st.error("The workbook could not be read. Check that it is a supported Excel file and layout.")
            with st.expander("Technical details"):
                st.code(f"{type(exc).__name__}: {exc}")

pending = st.session_state.pending_review
if pending is not None:
    missing = pending["missing"]
    for warning in pending["warnings"]:
        st.warning(warning)
    duplicate_groups = pending.get("duplicate_groups", [])
    if duplicate_groups:
        st.subheader("Review repeated lookup details")
        st.write(
            "These separate lookup rows share the same sample number, test period, "
            "and dilution. Review each group and choose whether workbook creation may continue."
        )
        decisions = []
        for index, group in enumerate(duplicate_groups, start=1):
            st.markdown(
                f"**Group {index}: Sample {group['sample_number']} · "
                f"{group['test_period']} · Dilution {group['dilution']}**"
            )
            for row in group["rows"]:
                st.write(f"Lookup label {row['label']} · Wells {', '.join(row['wells'])}")
                st.code(row["source_entry"], language=None)
            decision = st.radio(
                f"Decision for group {index}",
                ["Choose an action", "Continue with workbook creation", "Cancel workbook creation"],
                key=f"duplicate_decision_{pending['review_id']}_{index}",
                horizontal=True,
            )
            decisions.append(decision)
        if pending.get("duplicates_approved") and pending.get(
            "approved_duplicate_decisions"
        ) != decisions:
            pending.pop("duplicates_approved", None)
            pending.pop("approved_duplicate_decisions", None)
        if all(value == "Continue with workbook creation" for value in decisions):
            if st.button("Confirm review and continue", type="primary"):
                pending["duplicates_approved"] = True
                pending["approved_duplicate_decisions"] = decisions
                st.rerun()
        elif all(value != "Choose an action" for value in decisions) and any(
            value == "Cancel workbook creation" for value in decisions
        ):
            st.warning("At least one group is marked to cancel. No workbook will be created.")
            if st.button("Cancel workbook creation"):
                st.session_state.pending_review = None
                st.info("Review cancelled. No output workbook was created.")
        else:
            st.info("Choose an action for every flagged group before proceeding.")
    preview = pending.get("correction_preview")
    if preview is not None:
        st.subheader("Review corrections before creating the workbook")
        st.write("Check each original entry against the corrected entry. No workbook has been created yet.")
        corrections_by_well = preview["well_values"]
        wells = sorted({item["well"] for item in missing if item.get("well")})
        for well in wells:
            original = next(item["source_entry"] for item in missing if item.get("well") == well)
            corrections = corrections_by_well.get(well, {})
            corrected = corrections.get("Source metadata entry")
            if corrected is None:
                corrected = complete_source_entry(original, corrections)
            left, right = st.columns(2)
            with left:
                st.markdown(f"**Well {well} — original entry**")
                st.code(original or "(blank)", language=None)
            with right:
                st.markdown(f"**Well {well} — corrected entry**")
                st.code(corrected or "(blank)", language=None)
        if preview.get("study_override"):
            st.info(f"Study number to use in the output: {preview['study_override']}")

        create_col, edit_col, cancel_col = st.columns(3)
        create_now = create_col.button(
            "Create workbook with these corrections", type="primary",
            disabled=bool(duplicate_groups) and not pending.get("duplicates_approved"),
        )
        edit_now = edit_col.button("Edit corrections")
        cancel_now = cancel_col.button("Cancel and discard")
        if edit_now:
            pending.pop("correction_preview", None)
            pending.pop("duplicate_groups", None)
            pending.pop("duplicates_approved", None)
            pending.pop("approved_duplicate_decisions", None)
            st.rerun()
        elif cancel_now:
            st.session_state.pending_review = None
            st.session_state.pop(f"review_errors_{pending['review_id']}", None)
            st.info("Review cancelled. No output workbook was created.")
        elif create_now:
            try:
                with st.spinner("Creating the reformatted workbook…"):
                    with tempfile.TemporaryDirectory(prefix="plate_reformatter_") as work_dir:
                        source_path = Path(work_dir) / f"source{pending['suffix']}"
                        source_path.write_bytes(pending["source_bytes"])
                        if not pending.get("duplicates_approved"):
                            duplicate_groups = find_duplicate_lookup_groups(
                                source_path, well_values=corrections_by_well
                            )
                            if duplicate_groups:
                                pending["duplicate_groups"] = duplicate_groups
                                st.rerun()
                        output_path, study, sample_count, warnings = reformat(
                            source_path,
                            pending["plate_name"],
                            well_values=corrections_by_well,
                            study_override=preview.get("study_override", ""),
                            duplicate_groups_approved=bool(pending.get("duplicate_groups"))
                            and pending.get("duplicates_approved", False),
                        )
                        result_bytes = output_path.read_bytes()
                        output_name = output_path.name
                st.session_state.completed_result = {
                    "bytes": result_bytes,
                    "filename": output_name,
                    "study": study,
                    "sample_count": sample_count,
                    "warnings": warnings,
                }
                st.session_state.pending_review = None
                st.rerun()
            except Exception as exc:
                st.error("The workbook could not be created. No completed workbook is available.")
                with st.expander("Technical details"):
                    st.code(f"{type(exc).__name__}: {exc}")
    if missing:
        if preview is not None:
            missing = []
    if missing:
        st.warning("Required information is missing or needs correction. No output workbook has been created yet.")
        st.write("Fill in each item below using the source workbook, or cancel to stop.")
        errors_key = f"review_errors_{pending['review_id']}"
        validation_errors = st.session_state.get(errors_key, {})
        with st.form("complete_missing_information"):
            study_override = ""
            for item in missing:
                if item["kind"] == "study_number":
                    study_key = f"missing_study_number_{pending['review_id']}"
                    study_override = st.text_input(
                        "Study number — not found in the workbook",
                        key=study_key,
                        help="Enter the study number exactly as it should appear in the output.",
                    )
                    if study_key in validation_errors:
                        st.error(validation_errors[study_key])
                    continue
                if item["kind"] == "source_entry":
                    field_key = f"missing_{pending['review_id']}_{item['well']}_source_entry"
                    st.caption("Enter the complete corrected entry using this field order:")
                    st.code(
                        "[Study Number] [Custom ID] [Sample Number] [Group Number] "
                        "[Matrix/Split] [Test Period] [Dilution Fold]"
                    )
                    st.caption("Replace each bracketed name with its value; do not type the brackets.")
                    st.text_area(
                        f"Complete corrected source entry — well {item['well']}",
                        key=field_key,
                        height=90,
                        help="Use the displayed field order and verify every value against the source workbook.",
                    )
                    st.caption(f"Missing metadata: {item['reason']}")
                    st.caption(f"Original source entry: {item['source_entry']}")
                    if field_key in validation_errors:
                        st.error(validation_errors[field_key])
                    continue
                field_key = f"missing_{pending['review_id']}_{item['well']}_{item['field'].replace(' ', '_')}"
                field_help = (
                    f"Current source value: {item['value']}"
                    if item.get("value")
                    else f"Original source entry: {item['source_entry']}"
                )
                st.text_input(
                    f"{item['field']} — well {item['well']} (label {item['label']})",
                    key=field_key,
                    help=field_help,
                )
                if item.get("reason") != "Required value is missing.":
                    st.caption(item.get("reason", ""))
                if field_key in validation_errors:
                    st.error(validation_errors[field_key])
            complete = st.form_submit_button("Create workbook with these entries")

        cancel = st.button("Cancel and discard this review")
        if cancel:
            st.session_state.pending_review = None
            st.session_state.pop(errors_key, None)
            st.info("Review cancelled. No output workbook was created.")
        elif complete:
            well_values = {}
            new_errors = {}
            for item in missing:
                if item["kind"] == "study_number":
                    study_key = f"missing_study_number_{pending['review_id']}"
                    if not st.session_state.get(study_key, "").strip():
                        new_errors[study_key] = "Enter the study number to continue."
                    continue
                if item["kind"] == "source_entry":
                    key = f"missing_{pending['review_id']}_{item['well']}_source_entry"
                    value = st.session_state.get(key, "").strip()
                    if not value:
                        new_errors[key] = (
                            f"Enter the complete corrected source entry for well {item['well']}."
                        )
                    else:
                        issues = source_entry_issues(value)
                        if issues:
                            new_errors[key] = (
                                "The corrected entry still needs valid metadata: "
                                + ", ".join(issues) + "."
                            )
                        else:
                            well_values.setdefault(item["well"], {})[
                                "Source metadata entry"
                            ] = value
                    continue
                key = f"missing_{pending['review_id']}_{item['well']}_{item['field'].replace(' ', '_')}"
                value = st.session_state.get(key, "").strip()
                if not value:
                    new_errors[key] = f"Enter the missing {item['field'].lower()} for well {item['well']}."
                elif item["field"] == "Dilution Fold" and not is_numeric_value(value):
                    new_errors[key] = (
                        "Enter a number greater than zero, such as 5, 2.5, or 1,000. "
                        "Scientific notation is not accepted."
                    )
                else:
                    well_values.setdefault(item["well"], {})[item["field"]] = value
            if new_errors:
                st.session_state[errors_key] = new_errors
                st.rerun()
            else:
                st.session_state.pop(errors_key, None)
                study_key = f"missing_study_number_{pending['review_id']}"
                if any(item["kind"] == "study_number" for item in missing):
                    study_override = st.session_state.get(study_key, "").strip()
                pending["correction_preview"] = {
                    "well_values": well_values,
                    "study_override": study_override,
                }
                st.rerun()
    elif preview is None and (
        not duplicate_groups or pending.get("duplicates_approved")
    ):
        if duplicate_groups:
            st.success("All flagged lookup groups were approved. You can create the workbook now.")
        else:
            st.success("No required information is missing. You can create the workbook now.")
        if st.button(
            "Create workbook", type="primary",
            disabled=bool(duplicate_groups) and not pending.get("duplicates_approved"),
        ):
            try:
                with st.spinner("Creating the reformatted workbook…"):
                    with tempfile.TemporaryDirectory(prefix="plate_reformatter_") as work_dir:
                        source_path = Path(work_dir) / f"source{pending['suffix']}"
                        source_path.write_bytes(pending["source_bytes"])
                        if not pending.get("duplicates_approved"):
                            duplicate_groups = find_duplicate_lookup_groups(source_path)
                            if duplicate_groups:
                                pending["duplicate_groups"] = duplicate_groups
                                st.rerun()
                        output_path, study, sample_count, warnings = reformat(
                            source_path,
                            pending["plate_name"],
                            duplicate_groups_approved=bool(pending.get("duplicate_groups"))
                            and pending.get("duplicates_approved", False),
                        )
                        result_bytes = output_path.read_bytes()
                        output_name = output_path.name
                st.session_state.completed_result = {
                    "bytes": result_bytes,
                    "filename": output_name,
                    "study": study,
                    "sample_count": sample_count,
                    "warnings": warnings,
                }
                st.session_state.pending_review = None
                st.rerun()
            except Exception as exc:
                st.error("The workbook could not be created. No completed workbook is available.")
                with st.expander("Technical details"):
                    st.code(f"{type(exc).__name__}: {exc}")

result = st.session_state.completed_result
if result is not None:
    st.success(
        f"Reformat complete: {result['filename']} "
        f"(study {result['study']}; {result['sample_count']} grouped sample labels)."
    )
    if result["warnings"]:
        st.warning("Review these details against the source workbook:")
        for warning in result["warnings"]:
            st.markdown(f"- {warning}")

    st.subheader("Preview workbook contents")
    st.caption(
        "Review every output sheet below before downloading. The preview carries "
        "over the workbook's cell styles and uses a scrollable Excel-like grid."
    )
    try:
        preview_workbook = load_workbook(
            BytesIO(result["bytes"]), data_only=True, read_only=False
        )
        preview_tabs = st.tabs(preview_workbook.sheetnames)
        for preview_tab, sheet_name in zip(preview_tabs, preview_workbook.sheetnames):
            with preview_tab:
                worksheet = preview_workbook[sheet_name]
                st.markdown(_worksheet_preview_html(worksheet), unsafe_allow_html=True)
        preview_workbook.close()
    except Exception as exc:
        st.warning(
            "The workbook was created, but its contents could not be previewed. "
            "You can still download and review the file in Excel."
        )
        with st.expander("Preview technical details"):
            st.code(f"{type(exc).__name__}: {exc}")

    st.download_button(
        "Download reformatted workbook",
        data=result["bytes"],
        file_name=result["filename"],
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
    )

st.caption("Review the generated plate map. If the workbook includes Review Notes, review those details too.")
